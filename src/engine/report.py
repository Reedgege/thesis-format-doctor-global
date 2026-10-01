"""Comparison report: organize check results into a readable "which source wins" report.

Core value: explicitly list, for each dimension, how the four sources
(spec / school template / questionnaire / AI) are reconciled, so the user
can see at a glance what follows the general spec vs. the school requirement,
with conflicts highlighted.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Optional

from .checker import ALIGN_LABEL, Finding, summarize
from .docx_reader import DocumentProfile
from .questionnaire import TargetProfile
from .specs import get_spec

ORIGIN_LABEL = {
    "spec": "General Spec",
    "school": "School Template",
    "template": "School Template",
    "user": "Questionnaire",
    "ai_import": "AI Import",
    "ai": "AI Import",
    "latex_template": "LaTeX Template",
    "latex": "LaTeX Template",
}


@dataclass
class Report:
    meta: dict = field(default_factory=dict)
    resolution: list = field(default_factory=list)   # (dimension, source label)
    conflicts: list = field(default_factory=list)     # (dimension, left desc, right desc)
    findings: list = field(default_factory=list)
    summary: dict = field(default_factory=dict)
    notes: list = field(default_factory=list)         # questionnaire / LaTeX import notes
    markdown: str = ""


_DIM_LABEL = {
    "margin_top_in": "Top margin", "margin_bottom_in": "Bottom margin",
    "margin_left_in": "Left margin", "margin_right_in": "Right margin",
    "font_family": "Body font", "font_size_pt": "Body font size",
    "line_spacing": "Line spacing", "heading_levels": "Heading levels",
    "body_alignment": "Body alignment", "first_line_indent_in": "Body first-line indent",
    "space_after_pt": "Space after paragraph",
    "reference_style": "Reference format", "reference_hanging_indent_in": "Reference hanging indent",
    "reference_numbered": "Reference numbering",
}

_ALIGN_LABEL = ALIGN_LABEL  # single source of truth lives in checker


def _fmt_val(key: str, v) -> str:
    if v is None:
        return "—"
    if key in ("margin_top_in", "margin_bottom_in", "margin_left_in",
               "margin_right_in", "reference_hanging_indent_in",
               "first_line_indent_in"):
        return f'{v}"'
    if key == "font_size_pt":
        return f"{v}pt"
    if key == "line_spacing":
        return f"{v}x"
    if key == "space_after_pt":
        return f"{v}pt"
    if key == "body_alignment":
        return _ALIGN_LABEL.get(str(v), str(v))
    if key == "reference_numbered":
        return "Numbered" if v else "Author-Year"
    if key == "reference_style":
        return str(v)
    return str(v)


def _source_label(origin: str) -> str:
    return ORIGIN_LABEL.get(origin, origin)


def _friendly_source(target: TargetProfile) -> str:
    origins = sorted({v for v in target.resolution.values() if v})
    if not origins:
        return "Not set"
    return " + ".join(_source_label(o) for o in origins)


def _collect_conflicts(target: TargetProfile) -> list:
    """Conflicts between the general spec default and explicit sources, plus between explicit sources."""
    out: list = []
    spec_key = target.spec_key
    spec = get_spec(spec_key) if spec_key and spec_key != "Other" else None

    # 1) Spec default vs explicit source
    if spec:
        spec_page = spec.page.__dict__
        spec_ref = spec.reference
        for key, origin in target.resolution.items():
            if origin in ("spec",):
                continue
            if key in TargetProfile.PAGE_KEYS:
                tval = target.page.get(key)
                sval = spec_page.get(key)
            elif key == "heading_levels":
                tval = target.heading_levels
                sval = spec.heading_levels
            elif key == "reference_style":
                tval = target.reference_style
                sval = spec_ref.style_key
            elif key == "reference_hanging_indent_in":
                tval = target.reference_hanging_indent_in
                sval = spec_ref.hanging_indent_in
            elif key == "reference_numbered":
                tval = target.reference_numbered
                sval = spec_ref.numbered
            else:
                continue
            if tval is not None and sval is not None and tval != sval:
                out.append((_DIM_LABEL.get(key, key),
                            f"General Spec {_fmt_val(key, sval)}",
                            f"{_source_label(origin)} {_fmt_val(key, tval)}"))

    # 2) Between explicit sources
    for c in target.conflicts:
        dim = _DIM_LABEL.get(c["dim"], c["dim"])
        chain = " → ".join(
            f"{_source_label(lbl)} {_fmt_val(c['dim'], val)}" for lbl, val in c["entries"]
        )
        out.append((dim, "Source conflict", chain))

    return out


def build_report(docx_path: str, target: TargetProfile,
                 prof: DocumentProfile, findings: list[Finding]) -> Report:
    rep = Report()
    rep.meta = {
        "docx": docx_path,
        "source": _friendly_source(target),
        "spec_key": target.spec_key or "Other",
    }

    # —— Source resolution table ——
    for dim, origin in target.resolution.items():
        rep.resolution.append((_DIM_LABEL.get(dim, dim), _source_label(origin)))

    # —— Conflicts ——
    rep.conflicts = _collect_conflicts(target)

    rep.findings = findings
    rep.summary = summarize(findings)
    rep.notes = list(getattr(target, "notes", ()) or ())

    # —— Markdown ——
    lines: list = []
    lines.append("# PaperFormat Pro · Format Check Report")
    lines.append("")
    lines.append(f"- Document: `{docx_path}`")
    lines.append(f"- Spec: {rep.meta['spec_key']}  ·  Source: {rep.meta['source']}")
    s = rep.summary
    lines.append(f"- Summary: Pass {s.get('pass',0)} · Info {s.get('info',0)} · Warning {s.get('warn',0)} · Fail {s.get('fail',0)}")
    lines.append("")

    # Section numbers are assigned dynamically: only rendered sections take a number,
    # so we never jump from "1." to "3.".
    _sec = [0]

    def _sec_no() -> str:
        _sec[0] += 1
        return str(_sec[0])

    lines.append(f"## {_sec_no()}. Dimension Resolution (which source wins)")
    lines.append("")
    if rep.resolution:
        # Group by source: when a template is present the sources scatter;
        # grouping reads better than a flat list.
        groups: dict = {}
        for dim, origin in rep.resolution:
            groups.setdefault(origin, []).append(dim)
        for origin, dims in groups.items():
            lines.append(f"- **{origin}** ({len(dims)} items): {', '.join(dims)}")
    else:
        lines.append("- No specific target values set (upload a school template or fill the questionnaire)")
    lines.append("")
    lines.append("> Note: Page-level dimensions labeled \"General Spec\" (margins, fonts, line spacing, etc.) "
                 "are shown for reference only and do not force a Fail; page layout follows your school "
                 "template / questionnaire. Reference formatting always follows the general spec.")
    lines.append("")

    if getattr(target, "notes", None):
        lines.append(f"## {_sec_no()}. Notes")
        lines.append("")
        for n in target.notes:
            lines.append(f"- {n}")
        lines.append("")

    if rep.conflicts:
        lines.append(f"## {_sec_no()}. Spec vs School Template / Inter-source Conflicts")
        lines.append("")
        lines.append("> The following dimensions have conflicting requirements from different sources; "
                     "they were checked against the specific source (school / questionnaire / AI):")
        lines.append("")
        for dim, left, right in rep.conflicts:
            lines.append(f"- **{dim}**: {left} → {right}")
        lines.append("")

    lines.append(f"## {_sec_no()}. Detailed Findings")
    lines.append("")
    order = {"fail": 0, "warn": 1, "info": 2, "pass": 3}
    for f in sorted(findings, key=lambda x: order.get(x.severity, 9)):
        icon = {"pass": "✅", "warn": "⚠️", "fail": "❌", "info": "ℹ️"}.get(f.severity, "•")
        line = f"- {icon} **[{f.dimension}]** {f.message}"
        if f.expected or f.actual:
            line += f"  (Expected: {f.expected or '—'} | Actual: {f.actual or '—'})"
        lines.append(line)
    lines.append("")

    rep.markdown = "\n".join(lines)
    return rep


def build_docx_report(rep: "Report", out_path: str) -> str:
    """Render the structured report as a Word document (English).

    Mirrors the Markdown report's content and section order so the two
    outputs stay in lock-step. Returns the written path.

    The DOCX rendering deliberately avoids emoji (Word emoji glyphs vary by
    platform) and uses bracketed uppercase severity tags that read cleanly
    for Western users, e.g. ``[PASS] [margins] ...``.
    """
    from docx import Document

    doc = Document()

    # —— Title + meta block ——
    doc.add_heading("PaperFormat Pro · Format Check Report", level=0)
    meta = rep.meta
    s = rep.summary or {}
    doc.add_paragraph(f"Document: {meta.get('docx', '')}")
    doc.add_paragraph(f"Spec: {meta.get('spec_key', '')}  ·  Source: {meta.get('source', '')}")
    doc.add_paragraph(
        f"Summary: Pass {s.get('pass', 0)} · Info {s.get('info', 0)} · "
        f"Warning {s.get('warn', 0)} · Fail {s.get('fail', 0)}"
    )

    # Section numbers are assigned dynamically: only rendered sections take a number.
    _sec = [0]

    def _sec_no() -> int:
        _sec[0] += 1
        return _sec[0]

    # —— Section 1: Dimension Resolution ——
    doc.add_heading(f"{_sec_no()}. Dimension Resolution (which source wins)", level=1)
    if rep.resolution:
        groups: dict = {}
        for dim, origin in rep.resolution:
            groups.setdefault(origin, []).append(dim)
        for origin, dims in groups.items():
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{origin} ({len(dims)} items): ")
            r.bold = True
            p.add_run(", ".join(dims))
    else:
        doc.add_paragraph(
            "No specific target values set (upload a school template or fill the questionnaire)"
        )
    note_p = doc.add_paragraph()
    nr = note_p.add_run(
        "Note: Page-level dimensions labeled \"General Spec\" (margins, fonts, line spacing, etc.) "
        "are shown for reference only and do not force a Fail; page layout follows your school "
        "template / questionnaire. Reference formatting always follows the general spec."
    )
    nr.italic = True

    # —— Section 2: Notes ——
    if getattr(rep, "notes", None):
        doc.add_heading(f"{_sec_no()}. Notes", level=1)
        for n in rep.notes:
            doc.add_paragraph(n, style="List Bullet")

    # —— Section 3: Conflicts ——
    if rep.conflicts:
        doc.add_heading(f"{_sec_no()}. Spec vs School Template / Inter-source Conflicts", level=1)
        cp = doc.add_paragraph()
        cr = cp.add_run(
            "The following dimensions have conflicting requirements from different sources; "
            "they were checked against the specific source (school / questionnaire / AI):"
        )
        cr.italic = True
        for dim, left, right in rep.conflicts:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{dim}: ")
            r.bold = True
            p.add_run(f"{left} → {right}")

    # —— Section 4: Detailed Findings ——
    doc.add_heading(f"{_sec_no()}. Detailed Findings", level=1)
    order = {"fail": 0, "warn": 1, "info": 2, "pass": 3}
    sev_tag = {"pass": "PASS", "warn": "WARN", "fail": "FAIL", "info": "INFO"}
    for f in sorted(rep.findings, key=lambda x: order.get(x.severity, 9)):
        tag = sev_tag.get(f.severity, (f.severity or "").upper())
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"[{tag}] [{f.dimension}] ")
        r.bold = True
        p.add_run(f.message)
        if f.expected or f.actual:
            p.add_run(f"  (Expected: {f.expected or '—'} | Actual: {f.actual or '—'})")

    doc.save(out_path)
    return out_path

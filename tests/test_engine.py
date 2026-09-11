"""断言式回归测试套件（pytest）。

覆盖 codex 审查挑出的 P0/P1/P2 修复点：
- 固定行距不再误判 fail（P0-①）
- AI 填表 filled 包装不再静默丢弃（P0-②）
- 规范 vs 学校冲突进入报告（P0-③）
- 规范页面值仅作参考（info），不强制 fail（P1）
- reference_style 反填（P1）
- 损坏/非 docx 文档抛出 DocxReadError（P2）
- 多分节页边距不一致可检测（P2）
- 参考文献检查数据驱动（IEEE 编号 / APA 年份）
"""

from __future__ import annotations

import json
import os
import sys

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.engine import checker, report as report_mod
from src.engine.docx_reader import read_docx, DocxReadError
from src.engine.questionnaire import (
    target_from_spec, target_from_dict, target_from_ai_json,
    target_from_template_docx, merge_school, build_target,
)


def _make_doc(path, margin_in=1.0, font="Times New Roman", size_pt=12.0,
              spacing=2.0, ref_style="APA", spacing_rule=None):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(margin_in)
        s.bottom_margin = Inches(margin_in)
        s.left_margin = Inches(margin_in)
        s.right_margin = Inches(margin_in)
    p = doc.add_paragraph()
    run = p.add_run("This is a sample thesis body paragraph with some text to inspect.")
    run.font.name = font
    run.font.size = Pt(size_pt)
    pf = p.paragraph_format
    if spacing_rule == "exactly":
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(24)
    else:
        pf.line_spacing = spacing
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("Section", style="Heading 2")
    doc.add_paragraph("References")
    if ref_style == "APA":
        entry = doc.add_paragraph("Smith, J. A. (2020). Title of the work. Publisher.")
        hang = 0.5
    elif ref_style == "IEEE":
        entry = doc.add_paragraph('[1] J. A. Smith, "Title of the work," J. Name, vol. 1, no. 2, pp. 3-4, 2020.')
        hang = 0.0
    else:
        entry = doc.add_paragraph("Smith, J. A. (2020). Title of the work. Publisher.")
        hang = 0.0
    # 参考文献悬挂缩进与该风格规范保持一致（避免触发「悬挂缩进不达标」的独立检查）
    if hang:
        entry.paragraph_format.left_indent = Inches(hang)
        entry.paragraph_format.first_line_indent = Inches(-hang)
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# P0-① 固定行距
# ---------------------------------------------------------------------------

def test_fixed_line_spacing_not_false_fail(tmp_path):
    docx = _make_doc(tmp_path / "fixed.docx", spacing_rule="exactly")
    # 学校要求行距 2.0（严格，来源 school）→ 过去会把 24pt 读成 304800 误判 fail
    school = target_from_dict({
        "margin_top_in": 1.0, "margin_bottom_in": 1.0, "margin_left_in": 1.0,
        "margin_right_in": 1.0, "font_family": "Times New Roman",
        "font_size_pt": 12.0, "line_spacing": 2.0,
    }, source="school")
    prof, findings = checker.run_check(str(docx), school)
    ls = [f for f in findings if f.dimension == "line_spacing"]
    assert ls, "应有行距检查项"
    assert ls[0].severity != "fail", f"固定行距不应误判 fail，实际：{ls[0].severity}/{ls[0].actual}"
    assert prof.line_spacing is None, "固定值行距应读为 None（而非 EMU 倍数）"


# ---------------------------------------------------------------------------
# P0-② AI filled 包装
# ---------------------------------------------------------------------------

def test_ai_filled_wrapper_imported(tmp_path):
    jpath = tmp_path / "ai.json"
    jpath.write_text(json.dumps({"filled": {
        "margin_left_in": 1.5, "font_family": "Arial", "reference_style": "IEEE",
    }}), encoding="utf-8")
    t = target_from_ai_json(str(jpath))
    assert t.page.get("margin_left_in") == 1.5, "filled 包装的页边距应被读取"
    assert t.page.get("font_family") == "Arial"
    assert t.reference_style == "IEEE"


def test_ai_empty_raises(tmp_path):
    jpath = tmp_path / "empty.json"
    jpath.write_text(json.dumps({"foo": "bar"}), encoding="utf-8")
    try:
        target_from_ai_json(str(jpath))
        assert False, "空 JSON 应抛出 ValueError"
    except ValueError:
        pass


# ---------------------------------------------------------------------------
# P0-③ 冲突进入报告
# ---------------------------------------------------------------------------

def test_school_conflict_in_report(tmp_path):
    tpl = _make_doc(tmp_path / "tpl.docx", margin_in=1.5)  # 左页边距 1.5
    good = _make_doc(tmp_path / "good.docx", margin_in=1.0)
    target = build_target("APA", template=str(tpl))
    prof, findings = checker.run_check(str(good), target)
    rep = report_mod.build_report(str(good), target, prof, findings)
    # 左页边距应为学校来源硬判且 fail
    left = [f for f in findings if f.dimension == "margins" and "left" in f.message.lower()]
    assert left and left[0].severity == "fail" and left[0].source == "school"
    # 冲突应出现在报告
    assert any(d == "左页边距" for d, _, _ in rep.conflicts), "规范 vs 学校冲突应入报告"


# ---------------------------------------------------------------------------
# P1 规范页面值仅参考（info），不 fail
# ---------------------------------------------------------------------------

def test_spec_page_only_info_not_fail(tmp_path):
    good = _make_doc(tmp_path / "good.docx", margin_in=1.0)
    target = target_from_spec("APA")  # 纯规范，无学校
    prof, findings = checker.run_check(str(good), target)
    s = checker.summarize(findings)
    assert s["fail"] == 0, f"仅选规范时页面不应 fail，实际 {s}"
    # 引用维度仍应被检查（硬判，pass）
    ref = [f for f in findings if f.dimension == "reference"]
    assert ref and ref[0].severity in ("pass", "info")


def test_reference_backfill(tmp_path):
    t = target_from_dict({"reference_style": "IEEE"}, source="user")
    assert t.reference_numbered is True
    assert t.reference_hanging_indent_in == 0.0
    assert t.resolution.get("reference_numbered") == "spec"


# ---------------------------------------------------------------------------
# P2 健壮性
# ---------------------------------------------------------------------------

def test_garbage_file_raises(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a real docx file\x00\x01\x02")
    try:
        read_docx(str(bad))
        assert False, "损坏文件应抛 DocxReadError"
    except DocxReadError:
        pass


def test_multi_section_margins(tmp_path):
    doc = Document()
    doc.add_section()  # 现在有 2 个分节
    secs = doc.sections
    secs[0].left_margin = Inches(1.0)
    secs[1].left_margin = Inches(2.0)  # 第二节左页边距不同
    for s in secs:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
    doc.add_paragraph("Body text for reading.")
    path = tmp_path / "multi.docx"
    doc.save(str(path))
    prof = read_docx(str(path))
    assert prof.margins_inconsistent is True, "多分节页边距不一致应被标记"
    assert prof.margins.get("left") in (1.0, 2.0)


# ---------------------------------------------------------------------------
# P2 引用检查数据驱动
# ---------------------------------------------------------------------------

def test_ieee_numbered_pass(tmp_path):
    docx = _make_doc(tmp_path / "ieee.docx", ref_style="IEEE")
    target = target_from_spec("IEEE")
    prof, findings = checker.run_check(str(docx), target)
    ref = [f for f in findings if f.dimension == "reference"][0]
    assert ref.severity == "pass", f"IEEE 编号形态应 pass，实际 {ref.severity}"


def test_apa_missing_year_warn(tmp_path):
    # 构造一份参考文献缺 (Year) 的文档
    doc = Document()
    doc.add_paragraph("Body text for reading.")
    doc.add_paragraph("References")
    doc.add_paragraph("Smith, J. A. Title of the work. Publisher.")  # 缺 (Year)
    docx = tmp_path / "apa_noyear.docx"
    doc.save(str(docx))
    target = target_from_spec("APA")
    prof, findings = checker.run_check(str(docx), target)
    ref = [f for f in findings if f.dimension == "reference"][0]
    assert ref.severity == "warn", f"缺年份应 warn，实际 {ref.severity}"

"""海外版 v2.2.0 UI 重做回归测试（tkinter，真起窗口）。

为什么必须有这个文件：v2.2.0 把整层视觉重写了（圆角自绘卡片、圆角按钮、滚动区、
主按钮状态机），**同时**引入了几个只在"运行时"才暴露的坑 —— 它们全都不会让
pytest 里那堆纯逻辑用例变红，只有真把窗口建起来才能发现：

- `_run` 的结果整理段没有 `finally` → 抛异常后主按钮永久卡在 "Checking…"；
- 卡片内容变长（展开 Advanced options）时**不重新测量** → 新内容被圆角矩形裁掉，
  而且滚动条不出现、用户也滚不到；
- 次按钮 `set_visible` 后 pack 顺序被打乱 → 掉到辅助按钮**下面**；
- 换论文后旧结论不失效 → 主按钮还写着 "Fix N issues"，点下去改的是新文件。

这几条都是 Codex 2026-09-12 评审抓出来的真实缺陷，本文件把它们逐条钉住。

跑法：需要**带 tkinter** 的 Python。托管 venv 没有 tkinter 时整个文件自动 skip；
本机用系统 Python + venv 的 site-packages 可全量跑：

    PYTHONPATH=<venv>/Lib/site-packages python -m pytest tests/test_ui_redesign.py
"""

from __future__ import annotations

import importlib.util
import os
import sys
import tempfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)


# --------------------------------------------------------------------- 夹具
@pytest.fixture(scope="session")
def _tk_session():
    """整个会话共用一个 Tk root。

    同进程第二次 ``tk.Tk()`` 在本机系统 Python 上会因 Tcl 库加载失败抛 TclError
    （环境问题，非产品缺陷），共享 root 就绕开了，也让 GUI 回归能在 pytest 里跑。
    **这里刻意不 withdraw**：本文件要断言真实几何（页脚在可视区内、两栏等宽、
    右栏不随左栏展开被拽高），隐藏窗口拿不到这些；无法建窗的环境直接 skip。
    """
    if importlib.util.find_spec("tkinter") is None:
        pytest.skip("本机 Python 无 tkinter（GUI 回归需带 tkinter 的解释器）")
    import tkinter as tk

    try:
        root = tk.Tk()
    except Exception as e:                     # noqa: BLE001
        pytest.skip("Tk 不可用：%s" % e)
    yield root
    try:
        root.destroy()
    except Exception:
        pass


@pytest.fixture
def ui(tmp_path, monkeypatch, _tk_session):
    """建一个干净的 App（每个用例重建界面，避免状态串味）。"""
    monkeypatch.setenv("TFD_SETTINGS_FILE", str(tmp_path / "settings.json"))
    for w in list(_tk_session.winfo_children()):
        try:
            w.destroy()
        except Exception:
            pass
    from src.ui import app as ui_app

    a = ui_app.App(_tk_session)
    _tk_session.update_idletasks()
    yield a
    try:
        a._close_popups()
    except Exception:
        pass


def _make_docx(path, size_pt=9):
    """造一份"有明显格式问题"的 docx，保证体检能出 warn/fail。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("A Study of Something", level=1)
    run = doc.add_paragraph("Deliberately odd formatting.").runs[0]
    run.font.size = Pt(size_pt)
    doc.add_paragraph("Second paragraph.")
    doc.save(path)
    return path


def _mapped(root):
    """窗口是否真的映射到屏幕（无桌面会话时为 False）。"""
    try:
        return bool(root.winfo_ismapped())
    except Exception:
        return False


# --------------------------------------------------------------- 圆角自绘卡片
def test_cards_are_drawn_as_rounded_polygons(ui):
    """tkinter 的 Frame 只有直角，卡片圆角是 Canvas 自绘的 —— 必须真画出多边形。"""
    for card in (ui._left_card, ui._right_card):
        items = card._cv.find_withtag("card")
        assert items, "卡片没有画任何图形"
        assert "polygon" in {card._cv.type(i) for i in items}
        assert len(card._cv.coords(items[0])) // 2 > 8, "顶点太少，不是圆角"


# --------------------------------------------------------- P0-2 内容增长要重测
def test_advanced_toggle_remeasures_card_and_scrollregion(ui):
    """展开 Advanced options 后，卡片与滚动区都必须跟着长高。

    不重测的话新增的那两行被圆角矩形裁掉、滚动条也不出现 —— 用户看不到也滚不到。
    """
    from src.ui import theme

    before = ui._left_card._cv.winfo_reqheight()
    ui._toggle_advanced()
    ui.root.update_idletasks()
    need = ui._left_card.inner.winfo_reqheight() + 2 * theme.CARD_PAD_Y
    assert ui._left_card._cv.winfo_reqheight() >= need, "卡片请求高度没跟上"
    assert ui._left_card._cv.winfo_reqheight() > before, "展开前后请求高度没变化"

    region = str(ui._main_scroll._cv.cget("scrollregion")).split()
    assert int(region[3]) >= need, "scrollregion 没覆盖新增内容：%s" % region

    ui._toggle_advanced()
    ui.root.update_idletasks()
    assert not ui._adv_open


def test_collapse_toggles_visibility(ui):
    """winfo_manager() 判定 pack 状态：与窗口是否映射无关，CI 里也能断言。"""
    assert ui._adv_open is False
    assert ui._adv_body.winfo_manager() == ""
    ui._toggle_advanced()
    ui.root.update_idletasks()
    assert ui._adv_open is True
    assert ui._adv_body.winfo_manager() == "pack"
    ui._toggle_advanced()
    ui.root.update_idletasks()
    assert ui._adv_body.winfo_manager() == ""


# ------------------------------------------------------------ 主按钮状态机
def test_primary_button_state_machine(ui, tmp_path):
    """initial → ready → results → fixed 四态要走到位，且文案/可用态一致。"""
    from src.ui import i18n

    # initial（空态）：必填未满足，但主按钮仍可用（点一下给友好行内提示，v2.3.4）
    assert ui._check_btn.enabled is True
    assert ui._check_btn.text == i18n.t(ui.lang, "btn_check")
    assert ui._step_index == 0
    assert ui._fix_btn._visible is False
    assert "only required input" in ui.status_lbl.cget("text")
    # 空态不摆进度条（视觉收敛：避免右栏空荡时还挂着 Prepare→Check→Fix）
    assert ui._progress.winfo_manager() == ""

    # ready：选了真实存在的论文
    docx = _make_docx(tmp_path / "thesis.docx")
    ui.docx_path.set(docx)
    ui._refresh_rows()
    ui.root.update_idletasks()
    assert ui._phase == "ready"
    assert ui._check_btn.enabled is True
    assert ui._check_btn._command == ui._run
    assert ui._step_index == 1
    assert ui._fix_btn._visible is True
    # 选了论文后进度条重新出现
    assert ui._progress.winfo_manager() == "pack"

    # results：真跑一次体检
    ui._run()
    ui.root.update_idletasks()
    assert ui._phase == "results"
    assert isinstance(ui._last_issues, int)
    if ui._last_issues:
        assert ui._check_btn.text == i18n.t(ui.lang, "btn_fix_n", ui._last_issues)
        assert ui._check_btn._command == ui._run_fix
    assert ui._step_index == 2
    assert ui._report_text.strip(), "完整报告没留存，导不出来"

    # fixed：模拟落盘完成
    ui._last_output = str(tmp_path / "thesis_fixed.docx")
    open(ui._last_output, "wb").close()
    ui._phase = "fixed"
    ui._sync_ui()
    ui.root.update_idletasks()
    assert ui._check_btn.text == i18n.t(ui.lang, "btn_review")
    assert ui._step_index == 3, "三步应全部完成（进度条满格）"


def test_secondary_button_stays_above_aux_row(ui, tmp_path):
    """P1-1：次按钮 set_visible 后不能被排到辅助按钮行下面。"""
    docx = _make_docx(tmp_path / "t.docx")
    ui.docx_path.set(docx)
    ui._refresh_rows()
    ui.root.update_idletasks()
    assert ui._fix_btn._visible is True
    # winfo_rooty() 是绝对屏幕 Y：重设计把 _fix_btn 挪进了左卡、_aux_row 留在底部
    # CTA 条，两者父级不同，winfo_y()（父级相对）跨父比较没有意义；用绝对坐标才能
    # 正确断言"次按钮在辅助按钮行上方"。
    assert ui._fix_btn.winfo_rooty() < ui._aux_row.winfo_rooty()


# ---------------------------------------------------- P0-1 忙态必须能解除
def test_busy_state_is_released_when_result_processing_raises(ui, tmp_path, monkeypatch):
    """结果整理阶段抛异常时，`_busy` 必须被 finally 收回。

    否则主按钮永远停在 "Checking…"、次按钮永远隐藏 —— 只能重启进程。
    """
    from src.ui import app as ui_app
    from src.ui import i18n

    ui.docx_path.set(_make_docx(tmp_path / "t.docx"))
    ui._refresh_rows()

    class _Boom:
        @property
        def summary(self):
            raise RuntimeError("simulated result-processing failure")

        @property
        def markdown(self):
            raise RuntimeError("boom")

    monkeypatch.setattr(ui_app.report_mod, "build_report", lambda *a, **k: _Boom())
    ui._run()
    ui.root.update_idletasks()

    assert ui._busy is False
    assert ui._check_btn.enabled is True
    assert ui._check_btn.text != i18n.t(ui.lang, "btn_check_busy")
    assert ui._phase == "error"
    # 短行内提示 + 完整信息进可导出的报告（P2-3）
    assert "simulated result-processing failure" in ui.status_lbl.cget("text")
    assert "simulated result-processing failure" in ui._report_text


# ------------------------------------------------- P1-2 输入变更作废旧结论
def test_changing_document_invalidates_previous_result(ui, tmp_path):
    """换论文后旧结论必须作废，否则主按钮会拿着旧结论去改新文件。"""
    from src.ui import i18n

    ui.docx_path.set(_make_docx(tmp_path / "a.docx"))
    ui._refresh_rows()
    ui._run()
    ui.root.update_idletasks()
    assert ui._phase == "results"
    assert ui._checked_path.endswith("a.docx")

    ui.docx_path.set(_make_docx(tmp_path / "b.docx"))
    ui._refresh_rows()
    ui.root.update_idletasks()
    assert ui._phase == "ready"
    assert ui._last_issues is None
    assert ui._checked_path == ""
    assert ui._check_btn.text == i18n.t(ui.lang, "btn_check")


# ------------------------------------------------------ 语言往返 / 状态留存
def test_language_round_trip_rerenders_status_and_summary(ui, tmp_path):
    """切语言要按新语言重渲染状态行与摘要 —— 不重渲染就会出现空白或旧语言残留。"""
    from src.ui import i18n

    ui.docx_path.set(_make_docx(tmp_path / "t.docx"))
    ui._refresh_rows()
    ui._run()
    ui.root.update_idletasks()
    assert "Pass" in ui.summary_lbl.cget("text")
    en_status = ui.status_lbl.cget("text")
    assert en_status.strip()

    ui._set_lang("zh")
    ui._check_btn.draw()
    ui.root.update_idletasks()
    assert "通过" in ui.summary_lbl.cget("text"), "摘要没跟着切语言"
    assert ui.status_lbl.cget("text") != en_status, "状态行没跟着切语言"
    assert "问题" in ui.status_lbl.cget("text"), ui.status_lbl.cget("text")

    ui._set_lang("en")
    ui.root.update_idletasks()
    assert "Pass" in ui.summary_lbl.cget("text")
    assert ui.status_lbl.cget("text") == en_status


def test_english_ui_has_no_chinese_leak(ui):
    """英文界面不得出现中文（语言切换按钮按设计显示目标语言名除外）。"""
    import tkinter as tk

    def walk(w):
        yield w
        for c in w.winfo_children():
            yield from walk(c)

    texts = [str(w.cget("text")) for w in walk(ui.root) if isinstance(w, tk.Label)]
    lang_btn = ui.tr("lang_button")
    cjk = [t for t in texts
           if t != lang_btn and any("\u4e00" <= c <= "\u9fff" for c in t)]
    assert not cjk, "英文界面出现中文：%s" % cjk


# ------------------------------------------------------ 页脚 / 状态栏永久可见
def test_footer_and_statusbar_visible_within_window(ui):
    """v2.1.1 修过的 P0：页脚与状态栏绝不能被主体挤出可视区。

    无桌面会话拿不到真实几何 → 这条自动跳过（本机 GUI 冒烟覆盖）。
    """
    import tkinter as tk

    if not _mapped(ui.root):
        pytest.skip("窗口未映射到屏幕，拿不到真实几何")

    def walk(w):
        yield w
        for c in w.winfo_children():
            yield from walk(c)

    ui.root.geometry("1100x760")
    ui.root.update_idletasks()
    ui.root.update()
    win_h = ui.root.winfo_height()
    top = ui.root.winfo_rooty()

    found = {}
    for w in walk(ui.root):
        if not isinstance(w, tk.Label):
            continue
        t = str(w.cget("text"))
        if "hi@reedskill.com" in t:
            found["email"] = w
        elif "reedskill.com" in t:
            found["site"] = w
        elif "Thesis Format Doctor Global" in t:
            found["statusbar"] = w
    for name in ("email", "site", "statusbar"):
        assert name in found, "找不到 %s 标签" % name
        y = found[name].winfo_rooty() - top
        assert 0 < y < win_h, "%s 在可视区外（y=%d 窗口高=%d）" % (name, y, win_h)


def test_two_cards_equal_width_and_right_does_not_follow_left(ui):
    """两栏等宽；右栏按自身自然高度、贴顶，不随左栏展开 Advanced 被拽高。

    旧版强制两栏等高（sticky="nsew" + 行权重），左栏展开高级选项涨高时右栏也被
    拉下去（老板 2026-09-26 反馈"很难看"）。改为两卡 sticky="new" 后，右栏只取
    自身自然高度、不再被左栏拽高。这里钉死这条回归线。
    """
    if not _mapped(ui.root):
        pytest.skip("窗口未映射，拿不到真实几何")
    ui.root.update_idletasks()
    lw, lh = ui._left_card._cv.winfo_width(), ui._left_card._cv.winfo_height()
    rw, rh = ui._right_card._cv.winfo_width(), ui._right_card._cv.winfo_height()
    assert lw > 100 and lh > 100 and rh > 100
    # 等宽仍成立（列宽由窗口决定，与 sticky 无关）
    assert abs(lw - rw) <= 3, "两栏宽度不等：%d vs %d" % (lw, rw)
    # 右栏不再被强拉到与左栏等高：允许天然比左栏矮，但绝不能比左栏还高
    assert rh <= lh + 3, "右栏竟比左栏还高：%d vs %d" % (rh, lh)
    # 关键回归点：展开左栏 Advanced，右栏高度必须稳定（不跟着涨）
    right_before = ui._right_card._cv.winfo_height()
    left_before = ui._left_card._cv.winfo_height()
    ui._toggle_advanced()
    ui.root.update_idletasks()
    right_after = ui._right_card._cv.winfo_height()
    left_after = ui._left_card._cv.winfo_height()
    assert left_after > left_before + 20, \
        "展开 Advanced 后左栏没涨高：%d -> %d" % (left_before, left_after)
    assert abs(right_after - right_before) <= 3, \
        "展开 Advanced 后右栏被拽高：%d -> %d" % (right_before, right_after)
    ui._toggle_advanced()  # 还原，避免影响后续用例


# -------------------------------------------------- 切语言界面飘（v2.3.8 根治）
def test_main_top_gap_follows_header_even_when_scale_unchanged(ui):
    """切语言后窗口宽度不变 → scale 不变 → _on_resize 早退分支不得让主体顶留白
    停留在首帧的错误值（v2.3.7 切语言界面"飘"、最大化再还原才好的根因，v2.3.8 修）。

    机制：_top_gap() 依赖顶栏实测高度，重建/切语言首帧顶栏尚未测量（winfo_reqheight
    返回 1），旧代码只在 scale 变化时重设 pady，于是内容一直压在顶栏底下错位；只有
    真正的 resize 改了宽度才会纠正。修复后每次 _on_resize 都重设 pady。
    """
    if not _mapped(ui.root):
        pytest.skip("窗口未映射，拿不到真实几何")
    root = ui.root
    root.update_idletasks()

    def _top_pad():
        info = ui._main_scroll.pack_info().get("pady")
        return info[0] if isinstance(info, tuple) else info

    # 正常态：留白必须 == _top_gap()
    assert _top_pad() == ui._top_gap(), \
        "初始主体顶留白不对：%r != %r" % (_top_pad(), ui._top_gap())

    # 模拟切语言的早退场景：把 _last_scale 钉成与当前宽度一致，再触发一次 _on_resize
    ui._last_scale = ui.F.scale_for_width(root.winfo_width())
    ui._on_resize()
    root.update_idletasks()
    assert _top_pad() == ui._top_gap(), \
        "scale 不变时主体顶留白未被重设（早退 bug 回归）：%r != %r" \
        % (_top_pad(), ui._top_gap())

    # 顶栏底边必须落在主体顶部之上，不能重叠（切语言后内容不应压到顶栏下）
    hdr_bottom = (ui._hdr_brand.winfo_rooty() + ui._hdr_brand.winfo_height()
                  - root.winfo_rooty())
    assert hdr_bottom <= _top_pad() + 1, \
        "顶栏与主体重叠：hdr_bottom=%r top_pad=%r" % (hdr_bottom, _top_pad())


def test_set_lang_keeps_content_below_header(ui):
    """完整切语言路径：中↔英切换后内容顶留白仍跟随顶栏、不飘到顶栏下（v2.3.8）。"""
    if not _mapped(ui.root):
        pytest.skip("窗口未映射，拿不到真实几何")
    root = ui.root
    other = "zh" if ui.lang != "zh" else "en"
    ui._set_lang(other)
    root.update_idletasks()
    info = ui._main_scroll.pack_info().get("pady")
    top_pad = info[0] if isinstance(info, tuple) else info
    assert top_pad == ui._top_gap(), \
        "切语言后主体顶留白未跟随顶栏：%r != %r" % (top_pad, ui._top_gap())
    hdr_bottom = (ui._hdr_brand.winfo_rooty() + ui._hdr_brand.winfo_height()
                  - root.winfo_rooty())
    assert hdr_bottom <= top_pad + 1, \
        "切语言后顶栏与主体重叠：hdr_bottom=%r top_pad=%r" % (hdr_bottom, top_pad)


# ------------------------------------------------------------------ 滚动区
def test_scroll_area_shows_bar_only_when_needed(_tk_session):
    """滚动条只在内容超出时出现，且真的能滚、能回顶。"""
    import tkinter as tk

    from src.ui import theme
    from src.ui.widgets import ScrollArea

    win = tk.Toplevel(_tk_session)
    win.geometry("400x300")
    try:
        sa = ScrollArea(win, bg=theme.BG)
        sa.pack(fill="both", expand=True)
        for i in range(40):
            tk.Label(sa.inner, text="line %02d" % i, bg=theme.BG).pack(anchor="w")
        win.update_idletasks()
        win.update()
        assert sa._bar_shown is True, "内容超高却没出滚动条"

        before = sa._cv.yview()[0]
        sa._cv.yview_scroll(3, "units")
        assert sa._cv.yview()[0] > before, "滚不动"
        sa.scroll_to_top()
        assert sa._cv.yview()[0] <= 0.001
    finally:
        try:
            win.destroy()
        except Exception:
            pass

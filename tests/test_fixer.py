"""一键修正引擎断言式测试（pytest）。

守护铁律：一键修正只改格式、绝不改正文内容。
- 正文段落文本修正前后逐字一致；
- 页边距 / 字体 / 字号 / 行距按 TargetProfile 生效；
- 参考文献条目数不变、核心 token（作者/年份/题名）保留；
- 绝不覆盖原件（同路径必须报错）。
"""

from __future__ import annotations

import os
import sys

import pytest

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.engine.questionnaire import TargetProfile
from src.engine.fixer import fix_docx, preview_fix
from src.engine.docx_reader import reference_paragraph_indices


def _make_sample(path):
    doc = Document()
    for t in ("Body paragraph one.", "Body paragraph two with more words.", "Body paragraph three."):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
    doc.add_paragraph("References")
    doc.add_paragraph("Smith, J. (2020). Title of work. Publisher.")
    doc.add_paragraph("Lee, K. (2019). Another title. Journal.")
    for sec in doc.sections:
        sec.top_margin = Inches(2.0)
        sec.bottom_margin = Inches(2.0)
        sec.left_margin = Inches(2.0)
        sec.right_margin = Inches(2.0)
    doc.save(path)


def _body_texts(path):
    doc = Document(path)
    ref_idx = set(reference_paragraph_indices(doc))
    return [doc.paragraphs[i].text
            for i in range(len(doc.paragraphs)) if i not in ref_idx]


def _ref_texts(path):
    doc = Document(path)
    return [doc.paragraphs[i].text for i in reference_paragraph_indices(doc)]


def test_fix_preserves_body_text(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    before = _body_texts(src)
    target = TargetProfile(
        page={"margin_top_in": 1.25, "margin_bottom_in": 1.25,
              "margin_left_in": 1.0, "margin_right_in": 1.0,
              "font_family": "Arial", "font_size_pt": 11, "line_spacing": 1.5},
        reference_style="APA", reference_hanging_indent_in=0.5, reference_numbered=False,
    )
    fix_docx(str(src), str(out), target)
    assert _body_texts(out) == before


def test_fix_applies_margins(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    target = TargetProfile(page={"margin_top_in": 1.25, "margin_bottom_in": 1.25,
                                 "margin_left_in": 1.0, "margin_right_in": 1.0})
    fix_docx(str(src), str(out), target)
    doc = Document(str(out))
    for sec in doc.sections:
        assert sec.top_margin.inches == pytest.approx(1.25)
        assert sec.bottom_margin.inches == pytest.approx(1.25)
        assert sec.left_margin.inches == pytest.approx(1.0)
        assert sec.right_margin.inches == pytest.approx(1.0)


def test_fix_applies_font(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    target = TargetProfile(page={"font_family": "Arial", "font_size_pt": 11})
    fix_docx(str(src), str(out), target)
    doc = Document(str(out))
    for p in doc.paragraphs:
        for r in p.runs:
            if r.text.strip():
                assert r.font.name == "Arial"
                assert r.font.size == Pt(11)


def test_fix_applies_line_spacing(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    target = TargetProfile(page={"line_spacing": 1.5})
    fix_docx(str(src), str(out), target)
    doc = Document(str(out))
    for p in doc.paragraphs:
        assert p.paragraph_format.line_spacing == pytest.approx(1.5)
        # 必须是倍数制（命名 1.0/1.5/2.0 或 MULTIPLE），不能是固定值/最小值
        assert p.paragraph_format.line_spacing_rule in (
            WD_LINE_SPACING.MULTIPLE, WD_LINE_SPACING.SINGLE,
            WD_LINE_SPACING.ONE_POINT_FIVE, WD_LINE_SPACING.DOUBLE,
        )


def test_fix_reference_count_unchanged(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    before = reference_paragraph_indices(Document(str(src)))
    target = TargetProfile(reference_style="IEEE", reference_hanging_indent_in=0.0,
                           reference_numbered=True)
    fix_docx(str(src), str(out), target)
    after = reference_paragraph_indices(Document(str(out)))
    assert len(before) == len(after) == 2


def test_fix_reference_tokens_preserved_numbered(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    target = TargetProfile(reference_style="IEEE", reference_hanging_indent_in=0.0,
                           reference_numbered=True)
    fix_docx(str(src), str(out), target)
    joined = "\n".join(_ref_texts(str(out)))
    for tok in ("Smith", "2020", "Title of work", "Lee", "2019", "Another title"):
        assert tok in joined
    refs = _ref_texts(str(out))
    assert refs[0].startswith("[1]")
    assert refs[1].startswith("[2]")


def test_fix_reference_tokens_preserved_author_year(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_sample(src)
    target = TargetProfile(reference_style="APA", reference_hanging_indent_in=0.5,
                           reference_numbered=False)
    fix_docx(str(src), str(out), target)
    joined = "\n".join(_ref_texts(str(out)))
    for tok in ("Smith", "2020", "Title of work"):
        assert tok in joined
    # 非编号制不应加 [n] 前缀
    assert not _ref_texts(str(out))[0].lstrip().startswith("[")


def test_fix_writes_new_file_not_overwrite(tmp_path):
    src = tmp_path / "in.docx"
    _make_sample(src)
    target = TargetProfile(page={"font_family": "Arial", "font_size_pt": 11})
    with pytest.raises(ValueError):
        fix_docx(str(src), str(src), target)  # 同路径必须报错
    out = tmp_path / "out.docx"
    before = _body_texts(str(src))
    fix_docx(str(src), str(out), target)
    assert _body_texts(str(src)) == before  # 原件未被改
    assert out.exists()

"""端到端演示（直接运行 python tests/smoke_test.py 看报告）。

断言式回归请用 pytest 跑 tests/test_engine.py。
本脚本仅做「能跑通、报告可读」的肉眼演示。
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.shared import Pt, Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.engine import checker, report as report_mod
from src.engine.docx_reader import read_docx
from src.engine.questionnaire import build_target, target_from_template_docx


def make_doc(path, margin_in, font, size_pt, spacing, ref_style):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(margin_in)
        s.bottom_margin = Inches(margin_in)
        s.left_margin = Inches(margin_in)
        s.right_margin = Inches(margin_in)
    p = doc.add_paragraph()
    run = p.add_run("This is a sample thesis body paragraph with some text to inspect.")
    run.font.name = font
    run.font.size = Pt(size_pt)
    p.paragraph_format.line_spacing = spacing
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("Section", style="Heading 2")
    doc.add_paragraph("References")
    if ref_style == "APA":
        doc.add_paragraph("Smith, J. A. (2020). Title of the work. Publisher.")
    elif ref_style == "IEEE":
        doc.add_paragraph('[1] J. A. Smith, "Title of the work," J. Name, vol. 1, no. 2, pp. 3-4, 2020.')
    doc.save(path)
    return path


def run(label, docx, target):
    prof, findings = checker.run_check(docx, target)
    rep = report_mod.build_report(docx, target, prof, findings)
    print("=" * 60)
    print(label)
    print("=" * 60)
    print(rep.markdown)
    print()


def main():
    tmp = os.path.join(ROOT, "tests", "_samples")
    os.makedirs(tmp, exist_ok=True)

    good = make_doc(os.path.join(tmp, "good_apa.docx"), 1.0, "Times New Roman", 12.0, 2.0, "APA")
    bad = make_doc(os.path.join(tmp, "bad.docx"), 0.5, "Calibri", 10.0, 1.0, "APA")
    tpl = make_doc(os.path.join(tmp, "school_template.docx"), 1.5, "Times New Roman", 12.0, 2.0, "APA")

    # 1) 仅选规范（页面仅参考）
    run("样例1：仅选 APA（页面为参考，不应 fail）", good, build_target("APA"))

    # 2) 严格问卷 + 偏离文档（应 warn/fail）
    run("样例2：严格问卷 + 偏离文档（应 warn/fail）", bad,
        build_target("APA", questionnaire={
            "margin_top_in": 1.0, "margin_bottom_in": 1.0, "margin_left_in": 1.0,
            "margin_right_in": 1.0, "font_family": "Times New Roman",
            "font_size_pt": 12.0, "line_spacing": 2.0,
        }))

    # 3) 学校模板覆盖 + 冲突报告
    school = target_from_template_docx(read_docx(tpl), "APA")
    run("样例3：学校模板左页边距1.5 覆盖规范（冲突应入报告）", good,
        build_target("APA", template=tpl))

    print("冒烟演示完成。")


if __name__ == "__main__":
    main()

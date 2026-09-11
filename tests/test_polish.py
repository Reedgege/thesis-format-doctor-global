"""打磨轮新增能力的断言式测试（pytest）。

四组打磨对应的守护点：
① 补漏检维度：正文对齐 / 首行缩进 / 段后间距 / 参考文献悬挂缩进（检查 + 修正）
② 修正器风险：标题段恒被保护、输出不覆盖、预览只列实际会变的项
③ 准确度加固：表格文字纳入采样、参考文献识别不被正文里的字样误判
④ 体验工程：结构化 JSON 输出、choice 取值归一化
"""

from __future__ import annotations

import os
import sys

import pytest

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.engine import checker
from src.engine.docx_reader import read_docx, reference_paragraph_indices
from src.engine.fixer import fix_docx, preview_fix, unique_output_path
from src.engine.questionnaire import TargetProfile, target_from_dict, target_from_spec
from src.engine.specs import get_spec
from src.main import report_to_dict


# ---------------------------------------------------------------------------
# 样例构造工具
# ---------------------------------------------------------------------------

def _doc_with_body(path, alignment=None, indent_in=None, space_after_pt=None,
                   headings=False, refs=True, ref_hang_in=None,
                   body_font="Times New Roman", body_size=12.0,
                   line_spacing=None):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
    for t in ("First body paragraph.", "Second body paragraph.", "Third body paragraph."):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.font.name = body_font
        r.font.size = Pt(body_size)
        pf = p.paragraph_format
        if alignment is not None:
            pf.alignment = alignment
        if indent_in is not None:
            pf.first_line_indent = Inches(indent_in)
        if space_after_pt is not None:
            pf.space_after = Pt(space_after_pt)
        if line_spacing is not None:
            pf.line_spacing = line_spacing
    if headings:
        hp = doc.add_paragraph("Chapter One", style="Heading 1")
        hr = hp.add_run("Big Heading Text")
        hr.font.name = "Times New Roman"
        hr.font.size = Pt(18)
        hr.bold = True
    if refs:
        doc.add_paragraph("References")
        e = doc.add_paragraph("Smith, J. (2020). Title of work. Publisher.")
        if ref_hang_in:
            e.paragraph_format.left_indent = Inches(ref_hang_in)
            e.paragraph_format.first_line_indent = Inches(-ref_hang_in)
    doc.save(str(path))
    return str(path)


def _body_texts(path):
    doc = Document(str(path))
    ref_idx = set(reference_paragraph_indices(doc))
    return [doc.paragraphs[i].text
            for i in range(len(doc.paragraphs)) if i not in ref_idx]


# ---------------------------------------------------------------------------
# ① 补漏检维度：检查
# ---------------------------------------------------------------------------

def test_alignment_mismatch_from_school(tmp_path):
    docx = _doc_with_body(tmp_path / "a.docx", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    school = target_from_dict({
        "margin_top_in": 1.0, "margin_bottom_in": 1.0,
        "margin_left_in": 1.0, "margin_right_in": 1.0,
        "body_alignment": "justify",
    }, source="school")
    prof, findings = checker.run_check(docx, school)
    f = [x for x in findings if x.dimension == "alignment"]
    assert f, "应有正文对齐检查项"
    assert f[0].severity == "warn" and f[0].source == "school"
    assert prof.body_alignment == "left"


def test_alignment_spec_origin_downgraded(tmp_path):
    docx = _doc_with_body(tmp_path / "b.docx", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    # IEEE 规范要求两端对齐，但来源是 spec → 只提示，不报错
    prof, findings = checker.run_check(docx, target_from_spec("IEEE"))
    f = [x for x in findings if x.dimension == "alignment"][0]
    assert f.severity == "info", f"spec 来源的对齐不一致应降级 info，实际 {f.severity}"


def test_first_line_indent_checked(tmp_path):
    docx = _doc_with_body(tmp_path / "c.docx", indent_in=0.0)
    school = target_from_dict({
        "margin_top_in": 1.0, "first_line_indent_in": 0.5,
    }, source="school")
    prof, findings = checker.run_check(docx, school)
    f = [x for x in findings if x.dimension == "first_line_indent"][0]
    assert f.severity in ("warn", "fail")
    assert prof.first_line_indent_in == pytest.approx(0.0)


def test_space_after_checked(tmp_path):
    docx = _doc_with_body(tmp_path / "d.docx", space_after_pt=0.0)
    school = target_from_dict({
        "margin_top_in": 1.0, "space_after_pt": 12.0,
    }, source="school")
    prof, findings = checker.run_check(docx, school)
    f = [x for x in findings if x.dimension == "space_after"][0]
    assert f.severity in ("warn", "fail")
    assert prof.space_after_pt == pytest.approx(0.0)


def test_reference_hanging_indent_checked(tmp_path):
    # 参考文献条目没有悬挂缩进，但 APA 规范要求 0.5″ → 硬判不达标
    docx = _doc_with_body(tmp_path / "e.docx", ref_hang_in=None)
    prof, findings = checker.run_check(docx, target_from_spec("APA"))
    assert prof.reference_hanging_indent_in == pytest.approx(0.0)
    f = [x for x in findings if x.dimension == "reference_hanging_indent"]
    assert f and f[0].severity == "fail", "未设悬挂缩进应判不达标"


def test_reference_hanging_indent_pass_when_ok(tmp_path):
    docx = _doc_with_body(tmp_path / "e2.docx", ref_hang_in=0.5)
    prof, findings = checker.run_check(docx, target_from_spec("APA"))
    f = [x for x in findings if x.dimension == "reference_hanging_indent"][0]
    assert f.severity == "pass"


# ---------------------------------------------------------------------------
# ① 补漏检维度：修正
# ---------------------------------------------------------------------------

def test_fix_applies_body_layout(tmp_path):
    src = tmp_path / "f.docx"
    out = tmp_path / "f_out.docx"
    _doc_with_body(src, alignment=WD_ALIGN_PARAGRAPH.LEFT, indent_in=0.0)
    before = _body_texts(src)
    target = TargetProfile(page={"body_alignment": "justify",
                                 "first_line_indent_in": 0.5,
                                 "space_after_pt": 6.0})
    fix_docx(str(src), str(out), target)
    doc = Document(str(out))
    ref_idx = set(reference_paragraph_indices(doc))
    checked = 0
    for i, p in enumerate(doc.paragraphs):
        if i in ref_idx or not p.text.strip():
            continue
        checked += 1
        assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert p.paragraph_format.first_line_indent.inches == pytest.approx(0.5)
        assert p.paragraph_format.space_after.pt == pytest.approx(6.0)
    assert checked >= 3
    assert _body_texts(out) == before, "修正后正文文本必须逐字一致"


def test_fix_keeps_reference_hanging(tmp_path):
    src = tmp_path / "g.docx"
    out = tmp_path / "g_out.docx"
    _doc_with_body(src, indent_in=0.0, ref_hang_in=0.5)
    target = TargetProfile(page={"first_line_indent_in": 0.5},
                           reference_style="APA",
                           reference_hanging_indent_in=0.5,
                           reference_numbered=False)
    fix_docx(str(src), str(out), target)
    doc = Document(str(out))
    refs = [doc.paragraphs[i] for i in reference_paragraph_indices(doc)]
    assert refs, "应识别到参考文献条目"
    for p in refs:
        # 参考文献保持悬挂（负首行缩进），不被正文首行缩进覆盖
        assert p.paragraph_format.first_line_indent.inches == pytest.approx(-0.5)


# ---------------------------------------------------------------------------
# ② 修正器风险：标题保护 / 输出不覆盖 / 预览 diff
# ---------------------------------------------------------------------------

def test_fix_preserves_heading_format(tmp_path):
    src = tmp_path / "h.docx"
    out = tmp_path / "h_out.docx"
    _doc_with_body(src, headings=True)
    target = TargetProfile(page={"font_family": "Arial", "font_size_pt": 11,
                                 "line_spacing": 1.5, "body_alignment": "justify",
                                 "first_line_indent_in": 0.5,
                                 "space_after_pt": 6.0})
    fix_docx(str(src), str(out), target)
    doc = Document(str(out))
    heads = [p for p in doc.paragraphs if p.style.name.lower().startswith("heading")]
    assert heads, "样例应包含标题段"
    for hp in heads:
        for r in hp.runs:
            if r.text.strip():
                assert r.font.size != Pt(11), f"标题字号被压成正文了：{r.font.size}"
                assert r.font.name != "Arial", f"标题字体被改动了：{r.font.name}"
        assert hp.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
        assert hp.paragraph_format.line_spacing is None
        assert hp.paragraph_format.first_line_indent is None
        hr = [r for r in hp.runs if r.text == "Big Heading Text"][0]
        assert hr.font.size == Pt(18), "标题原有字号必须原样保留"
    # 正文仍按目标生效
    body = [p for p in doc.paragraphs
            if p.text.startswith("First body")][0]
    assert body.runs[0].font.size == Pt(11)


def test_unique_output_path(tmp_path):
    p = tmp_path / "x_fixed.docx"
    assert unique_output_path(str(p)) == str(p)
    p.write_text("occupied", encoding="utf-8")
    assert unique_output_path(str(p)) == str(tmp_path / "x_fixed-2.docx")
    (tmp_path / "x_fixed-2.docx").write_text("occupied", encoding="utf-8")
    assert unique_output_path(str(p)) == str(tmp_path / "x_fixed-3.docx")


def test_preview_lists_only_real_changes(tmp_path):
    src = tmp_path / "p.docx"
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
    p = doc.add_paragraph("Body text.")
    p.paragraph_format.line_spacing = 1.0
    doc.save(str(src))
    target = TargetProfile(page={"margin_top_in": 1.0, "line_spacing": 2.0})
    text = "\n".join(preview_fix(str(src), target))
    assert "页边距" not in text, "已符合的项不应出现在预览里"
    assert "行距" in text and "→" in text, "行距变化应显示为「现状 → 目标」"


# ---------------------------------------------------------------------------
# ③ 准确度加固
# ---------------------------------------------------------------------------

def test_table_text_sampled(tmp_path):
    path = tmp_path / "t.docx"
    doc = Document()
    bp = doc.add_paragraph()
    br = bp.add_run("Ordinary body paragraph.")
    br.font.name = "Times New Roman"
    br.font.size = Pt(12)
    tbl = doc.add_table(rows=1, cols=2)
    for cell in tbl.rows[0].cells:
        r = cell.paragraphs[0].add_run("Cell text")
        r.font.name = "Calibri"
        r.font.size = Pt(9)
    doc.save(str(path))
    prof = read_docx(str(path))
    assert prof.table_paragraph_count >= 2, "表格单元格段落应被采样"
    assert prof.font_family == "Calibri", "表格文字应计入字体统计"
    assert prof.font_size_pt == pytest.approx(9.0)


def test_reference_detection_ignores_body_mention(tmp_path):
    path = tmp_path / "r.docx"
    doc = Document()
    doc.add_paragraph("This chapter reviews the method.")
    doc.add_paragraph("See the References section for details.")
    doc.add_paragraph("Another ordinary body paragraph follows.")
    doc.save(str(path))
    prof = read_docx(str(path))
    assert prof.reference_entries == [], "正文里出现的 References 字样不应被当成参考文献段"


def test_reference_detection_still_works(tmp_path):
    path = tmp_path / "r2.docx"
    doc = Document()
    doc.add_paragraph("Body text before the list.")
    doc.add_paragraph("References")
    doc.add_paragraph("Smith, J. (2020). Title of work. Publisher.")
    doc.save(str(path))
    prof = read_docx(str(path))
    assert len(prof.reference_entries) == 1


# ---------------------------------------------------------------------------
# ④ 体验工程
# ---------------------------------------------------------------------------

def test_report_dict_structure(tmp_path):
    src = _doc_with_body(tmp_path / "j.docx", ref_hang_in=0.5)
    target = target_from_spec("APA")
    prof, findings = checker.run_check(src, target)
    d = report_to_dict(src, target, prof, findings)
    for k in ("meta", "summary", "resolution", "conflicts", "findings", "profile"):
        assert k in d
    for k in ("body_alignment", "first_line_indent_in", "space_after_pt",
              "table_paragraph_count", "reference_hanging_indent_in"):
        assert k in d["profile"], f"结构化输出缺少 {k}"


def test_choice_value_normalized():
    t = target_from_dict({"margin_top_in": 1.0, "body_alignment": "JUSTIFY"},
                         source="school")
    assert t.page.get("body_alignment") == "justify"
    t2 = target_from_dict({"margin_top_in": 1.0, "body_alignment": "middle"},
                          source="school")
    assert "body_alignment" not in t2.page, "不在白名单里的取值应丢弃而非臆造"


def test_spec_body_layout_defaults():
    apa = get_spec("APA")
    assert apa.page.body_alignment == "left"
    assert apa.page.first_line_indent_in == pytest.approx(0.5)
    assert apa.page.space_after_pt == pytest.approx(0.0)
    ieee = get_spec("IEEE")
    assert ieee.page.body_alignment == "justify"
    assert ieee.page.first_line_indent_in == pytest.approx(0.0)
    other = get_spec("Other")
    assert other.page.body_alignment is None

"""CLI 端到端冒烟：真实 subprocess + 真实 docx，跑完打印结论。

为什么要有它：`pytest` 主要覆盖单元/组件级行为，而**授权门禁与入参校验的执行顺序**、
「原件字节未变」、「rc=0/2/3 三条路径」这类只有在真进程里才暴露的问题，单测挡不住
（本轮就靠它抓到「`--out` 指向原件被报成试用已用完」等 4 个问题）。

用法（Windows）::

    .venv\\Scripts\\python.exe tools\\smoke_cli.py
    # 或指定输出文件（推荐，避免中文控制台 GBK 编码问题）
    .venv\\Scripts\\python.exe tools\\smoke_cli.py -o smoke_out.txt

约定：脚本自身**不往控制台 print 中文/emoji**（Windows GBK 控制台会 UnicodeEncodeError），
统一收集后写 UTF-8 文件；末尾以 `FAILED` 或 `ALL OK` 收尾，退出码 0/1 可用于 CI。
"""

from __future__ import annotations

import argparse
import io
import json
import os
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PY = sys.executable

# 期望值（与 docs/需求规格与功能说明.md 的退出码约定一致）
RC_OK, RC_USAGE, RC_GATE = 0, 2, 3


class Smoke:
    def __init__(self) -> None:
        self.buf = io.StringIO()
        self.failures: list[str] = []

    def say(self, *a) -> None:
        self.buf.write(" ".join(str(x) for x in a) + "\n")

    def check(self, label: str, cond: bool, detail: str = "") -> None:
        mark = "ok  " if cond else "FAIL"
        if not cond:
            self.failures.append(label + (f" :: {detail}" if detail else ""))
        self.say(f"   [{mark}] {label}" + (f"  ({detail})" if detail else ""))


def _run(args, env, cwd=ROOT):
    return subprocess.run([PY, "-m", "src.main"] + args, cwd=cwd, env=env,
                          capture_output=True)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--out", default=None,
                    help="把完整报告写入该文件（UTF-8）；默认写到临时目录")
    args = ap.parse_args()

    from docx import Document
    from docx.shared import Inches, Pt

    s = Smoke()
    tmp = tempfile.mkdtemp(prefix="tfd_smoke_")

    def env_for(name="license.json"):
        e = os.environ.copy()
        e["PYTHONPATH"] = ROOT
        e["TFD_LICENSE_FILE"] = os.path.join(tmp, name)
        return e

    env = env_for()

    # ---- 0. 造样本：字体/字号/页边距/首行缩进都不合规，且参考文献无悬挂缩进
    src = os.path.join(tmp, "sample.docx")
    doc = Document()
    for text in ("Intro paragraph of the thesis.", "Methods paragraph.",
                 "Conclusion paragraph."):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
    doc.add_paragraph("References")
    doc.add_paragraph("Smith, J. (2020). Title of work. Publisher.")
    for sec in doc.sections:
        sec.top_margin = Inches(2.0)
        sec.left_margin = Inches(2.0)
    doc.save(src)
    src_bytes = open(src, "rb").read()
    s.say(f"[0] 样本已生成：{src}")

    # ---- 1. check：人类可读报告，章节号连续、数值无重复前缀
    r = _run(["check", src, "--spec", "APA"], env)
    s.say("\n[1] check（人类可读）")
    md = r.stdout.decode("utf-8", "replace")
    s.check("rc=0", r.returncode == RC_OK, f"rc={r.returncode}")
    heads = [ln for ln in md.splitlines() if ln.startswith("## ")]
    s.check("章节号连续（一、二…）",
            bool(heads) and all(h.startswith(f"## {'一二三四五六七八九十'[i]}、")
                                for i, h in enumerate(heads)), str(heads))
    s.check("无「一之二」", "一之二" not in md)
    s.check("无「实测：实测」", "实测：实测" not in md)

    # ---- 2. check --json：合法 UTF-8 可解析
    r = _run(["check", src, "--spec", "APA", "--json"], env)
    s.say("\n[2] check --json")
    try:
        payload = json.loads(r.stdout.decode("utf-8"))
        s.check("UTF-8 可解析且含 findings",
                bool(payload.get("findings")), f"{len(payload.get('findings', []))} 条")
    except Exception as exc:  # noqa: BLE001
        s.check("UTF-8 可解析", False, repr(exc))

    # ---- 3. fix --preview：不落盘
    r = _run(["fix", src, "--spec", "APA", "--preview"], env)
    s.say("\n[3] fix --preview")
    s.check("rc=0", r.returncode == RC_OK, f"rc={r.returncode}")
    s.check("未落盘", not os.path.exists(os.path.join(tmp, "sample_fixed.docx")))

    # ---- 4. fix 第一次：rc=0，提示已用尽，原件不变
    r = _run(["fix", src, "--spec", "APA"], env)
    s.say("\n[4] fix 第一次（首次免费试用）")
    out4 = r.stdout.decode("utf-8", "replace")
    fixed = os.path.join(tmp, "sample_fixed.docx")
    s.check("rc=0", r.returncode == RC_OK, f"rc={r.returncode}")
    s.check("生成修正稿", os.path.exists(fixed))
    s.check("提示已用尽（非「首次免费试用」）",
            "已用尽" in out4 and "首次免费试用" not in out4)
    s.check("原件字节未变", open(src, "rb").read() == src_bytes)
    if os.path.exists(fixed):
        d2 = Document(fixed)
        s.check("左页边距已改为 1.0\"",
                abs(d2.sections[0].left_margin.inches - 1.0) < 1e-6)
        s.check("正文文字未被改动",
                d2.paragraphs[0].text == "Intro paragraph of the thesis.")

    # ---- 5. fix 第二次：试用耗尽 → rc=3
    r = _run(["fix", src, "--spec", "APA"], env)
    s.say("\n[5] fix 第二次（试用已用尽）")
    s.check("rc=3", r.returncode == RC_GATE, f"rc={r.returncode}")

    # ---- 6. 关键回归：--out 指向原件 → 必须先于门禁返回 rc=2
    r = _run(["fix", src, "--spec", "APA", "--out", os.path.join(tmp, "SAMPLE.DOCX")], env)
    s.say("\n[6] --out 指向原件（大小写变体，试用已用尽）")
    s.check("rc=2（参数错误不报成试用问题）",
            r.returncode == RC_USAGE, f"rc={r.returncode}")
    s.check("原件字节未变", open(src, "rb").read() == src_bytes)

    # ---- 7. 关键回归：已合规文档不触发门禁、不落盘
    r = _run(["fix", fixed, "--spec", "APA"], env)
    s.say("\n[7] 已合规文档（试用已用尽）")
    s.check("rc=0", r.returncode == RC_OK, f"rc={r.returncode}")
    s.check("提示已符合、未消耗额度",
            "已符合目标格式" in r.stdout.decode("utf-8", "replace"))
    s.check("未生成新副本",
            not os.path.exists(os.path.join(tmp, "sample_fixed_fixed.docx")))

    # ---- 8. status
    r = _run(["status"], env)
    s.say("\n[8] status")
    s.check("rc=0", r.returncode == RC_OK, f"rc={r.returncode}")

    # ---- 9. 损坏文档：友好报错，不 traceback
    bad = os.path.join(tmp, "bad.docx")
    with open(bad, "wb") as f:
        f.write(b"not a docx at all")
    r = _run(["check", bad, "--spec", "APA"], env)
    s.say("\n[9] 损坏文档")
    s.check("rc=2", r.returncode == RC_USAGE, f"rc={r.returncode}")
    err = r.stderr.decode("utf-8", "replace")
    s.check("友好报错且无 traceback", "Traceback" not in err and "无法读取" in err)

    # ---- 10. 修正稿再体检：不达标应为 0
    r = _run(["check", fixed, "--spec", "APA"], env)
    s.say("\n[10] 修正稿再体检")
    md10 = r.stdout.decode("utf-8", "replace")
    conc = [ln for ln in md10.splitlines() if ln.startswith("- 结论")]
    s.check("rc=0", r.returncode == RC_OK, f"rc={r.returncode}")
    s.check("不达标 0", bool(conc) and "不达标 0" in conc[0], conc[0] if conc else "")

    # ---- 11. 全新账号 + --out 原件：rc=2 且不消耗额度
    fresh = env_for("lic2.json")
    r = _run(["fix", src, "--spec", "APA", "--out", os.path.join(tmp, "SAMPLE.DOCX")], fresh)
    s.say("\n[11] 全新账号 + --out 原件")
    s.check("rc=2", r.returncode == RC_USAGE, f"rc={r.returncode}")
    s.check("未消耗试用（无状态文件）",
            not os.path.exists(os.path.join(tmp, "lic2.json")))

    # ---- 收尾
    s.say("")
    if s.failures:
        s.say(f"FAILED（{len(s.failures)} 项）：")
        for f in s.failures:
            s.say("  - " + f)
    else:
        s.say("ALL OK —— CLI 端到端 11 组检查全部通过。")
    s.say(f"\n临时目录：{tmp}")

    out_path = args.out or os.path.join(tempfile.gettempdir(), "tfd_smoke_out.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(s.buf.getvalue())
    print(out_path)
    return 1 if s.failures else 0


if __name__ == "__main__":
    sys.exit(main())
